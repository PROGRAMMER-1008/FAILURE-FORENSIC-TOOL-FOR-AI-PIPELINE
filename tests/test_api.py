import asyncio
import io
import httpx
import pytest
from docx import Document
from ffv2.api import create_app
from ffv2.settings import Settings
from ffv2.uploads import DocumentError, extract
from ffv2.worker import SYNTHETIC_SOURCE, run_one


def request(app, method, path, **kwargs):
    async def call():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://testserver"
        ) as client:
            return await client.request(method, path, **kwargs)

    # Windows asyncio creates a socket pair for its loop, without network traffic.
    return asyncio.run(call())


@pytest.mark.enable_socket
def test_api_workflow_and_contracts(store):
    app = create_app(Settings(), store)
    body = {"text": SYNTHETIC_SOURCE, "injected": True, "idempotency_key": "run-key"}
    response = request(app, "POST", "/api/v1/runs", json=body)
    assert response.status_code == 202
    run_id = response.json()["id"]
    assert request(app, "POST", "/api/v1/runs", json=body).json()["id"] == run_id
    assert request(app, "GET", "/api/v1/runs/not-a-uuid").status_code == 422
    assert (
        request(app, "GET", "/api/v1/runs/00000000-0000-0000-0000-000000000000").status_code == 404
    )
    run_one(store, Settings())
    run = request(app, "GET", f"/api/v1/runs/{run_id}").json()
    review = {
        "idempotency_key": "review-key",
        "reviewer": "Local tester",
        "revision": run["revision"],
        "target_finding_revision": 1,
        "verdict": "correct",
        "origins": ["extract"],
    }
    assert request(app, "POST", f"/api/v1/runs/{run_id}/reviews", json=review).status_code == 200
    assert request(app, "POST", f"/api/v1/runs/{run_id}/reviews", json=review).status_code == 200
    assert len(request(app, "GET", "/api/v1/datasets").json()) == 1
    events = request(app, "GET", f"/api/v1/runs/{run_id}/events", headers={"Last-Event-ID": "2"})
    assert "id: 1\n" not in events.text and "event: succeeded" in events.text


@pytest.mark.enable_socket
def test_demo_mutation_and_upload_limit(store):
    demo = create_app(Settings(mode="demo"), store)
    assert request(demo, "POST", "/api/v1/runs", json={}).status_code == 403
    app = create_app(Settings(max_upload_bytes=20), store)
    assert request(app, "POST", "/api/v1/uploads", content=b"a" * 21).status_code == 413


def test_table_only_docx_and_bad_files():
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic table invoice"
    data = io.BytesIO()
    document.save(data)
    assert "Synthetic table invoice" in extract(
        data.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    for data, mime in [
        (b"bad", "application/pdf"),
        (b"\xff", "text/plain"),
        (b"bad", "application/zip"),
    ]:
        with pytest.raises(DocumentError):
            extract(data, mime)
